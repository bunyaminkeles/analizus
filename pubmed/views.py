import re
from functools import wraps

from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, Http404, HttpResponse
from django.shortcuts import render, get_object_or_404
from django.utils.translation import gettext, pgettext
from django.views.decorators.http import require_GET, require_POST
from django_ratelimit.decorators import ratelimit

from .forms import PubMedSearchForm, FIELD_CHOICES
from .models import PubMedSearchJob
from .services.job_runner import run_scraping_job


def _safe_filename(job, ext: str) -> str:
    """pubmed_<arama_kelimesi>_<YYYYMMDD>.<ext> formatında güvenli dosya adı."""
    from django.utils import timezone
    keyword = next((p.get('value', '').strip() for p in job.query_parts if p.get('value', '').strip()), '')
    keyword = re.sub(r'[^\w\s-]', '', (keyword or 'sonuclar')[:40], flags=re.UNICODE)
    keyword = re.sub(r'\s+', '_', keyword).strip('_')
    return f'pubmed_{keyword}_{timezone.now():%Y%m%d}.{ext}'


def feature_required(flag_name):
    """Decorator: SiteSettings'deki feature flag kapalıysa 404 döner."""
    def decorator(view_func):
        @wraps(view_func)
        def wrapper(request, *args, **kwargs):
            from forum.models import SiteSettings
            if not getattr(SiteSettings.load(), f'feature_{flag_name}', False):
                raise Http404
            return view_func(request, *args, **kwargs)
        return wrapper
    return decorator


def _is_ajax(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest'


@feature_required('pubmed')
@ratelimit(key='user', rate='10/h', method='POST', block=True)
def pubmed_landing(request):
    """Arama sayfası (giriş yapmamışa tanıtım). Arama AJAX ile başlatılır, sonuç polling ile gelir."""
    if not request.user.is_authenticated:
        return render(request, 'service_promo.html', {
            'promo_title': gettext('PubMed Yayın Kazıma ve Veri İndirme Aracı'),
            'promo_icon': 'bi-heart-pulse',
            'promo_color': 'primary',
            'promo_description': gettext('40 milyondan fazla biyomedikal yayından kodsuz veri kazıma aracı. Başlık, özet, yazar, MeSH terimi ve kurum alanlarında arama yapın; sonuçları Excel veya TXT olarak indirin.'),
            'promo_features': [
                {'icon': 'bi-download', 'title': gettext('Excel & TXT İndirme'), 'desc': gettext('Başlık, yazar, dergi, yıl, DOI, PMID, MeSH terimleri ve özet verilerini Excel veya TXT olarak indirin.')},
                {'icon': 'bi-database-fill', 'title': gettext('40M+ Biyomedikal Kayıt'), 'desc': gettext("PubMed'in tıp, hemşirelik, sağlık ve yaşam bilimleri veri tabanında kodsuz arama yapın.")},
                {'icon': 'bi-sliders', 'title': gettext('Kodsuz Gelişmiş Sorgulama'), 'desc': gettext('Başlık, özet, yazar, MeSH terimi, dergi, kurum ve yıl alanlarını birleştirerek arama yapın.')},
                {'icon': 'bi-bar-chart-line-fill', 'title': gettext('Tek Tıkla Bibliometrik Analiz'), 'desc': gettext('100+ sonuçta tek tıkla bibliometrik analize gönderin, PDF rapor alın.')},
            ],
            'promo_steps': [
                gettext('Arama kriterlerinizi seçin: başlık, MeSH terimi, yazar veya yıl.'),
                gettext("Sistem PubMed'den verileri kazıyarak saniyeler içinde listeler."),
                gettext('Sonuçları Excel veya TXT olarak indirin ya da e-postanıza gönderin.'),
            ],
        })

    user = request.user
    daily_limit = PubMedSearchJob.get_daily_limit(user)
    remaining = max(0, daily_limit - PubMedSearchJob.daily_count_for_user(user))

    if request.method == 'POST':
        if not _is_ajax(request):
            raise Http404
        form = PubMedSearchForm(request.POST)
        if not form.is_valid():
            errors = form.errors.get('query_parts_json', [gettext('Geçersiz sorgu.')])
            return JsonResponse({'error': errors[0]}, status=400)
        if not user.profile.email_verified:
            return JsonResponse({'error': gettext('PubMed tarama için e-posta doğrulaması gereklidir. Profil sayfanızdan e-postanızı doğrulayın.')}, status=403)
        if remaining <= 0:
            # Tam cümle msgid'ler (OpenAlex ile aynı) — premium olmayana 7 aramalık premium önerilir
            msg = (gettext('Günlük demo limitiniz doldu ({used}/{limit}). Premium üyelikle 7 aramaya yükseltebilirsiniz.')
                   if not user.profile.is_premium else
                   gettext('Günlük demo limitiniz doldu ({used}/{limit}). Yarın tekrar deneyebilirsiniz.'))
            return JsonResponse({'error': msg.format(used=daily_limit, limit=daily_limit)}, status=429)

        job = PubMedSearchJob.objects.create(user=user, query_parts=form.cleaned_data['query_parts_json'])
        run_scraping_job(job.id)
        return JsonResponse({'status': 'started', 'job_id': str(job.id), 'remaining': remaining - 1})

    active_job = PubMedSearchJob.objects.filter(user=user, status__in=['pending', 'running']).order_by('-created_at').first()
    return render(request, 'pubmed/landing.html', {
        'form': PubMedSearchForm(),
        'remaining': remaining,
        'daily_limit': daily_limit,
        'active_job_id': str(active_job.id) if active_job else None,
        # JS'e json_script ile — {ad} yer tutucular JS'te doldurulur (tam cümle msgid)
        'js_config': {
            'fields': [[value, str(label)] for value, label in FIELD_CHOICES],
            'texts': {
                'need_criteria': gettext('Lütfen en az bir arama kriteri girin.'),
                'searching': gettext('Aranıyor...'),
                'queue': gettext('Sırada bekliyorsunuz — tahmini sıra: {pos}.'),
                'starting': gettext('Sıradaki tarama başlıyor...'),
                'running': gettext("PubMed'den veriler çekiliyor, lütfen bekleyin..."),
                'resumed': gettext('Aramanız arka planda devam ediyor, lütfen bekleyin...'),
                'failed': gettext('Arama başarısız oldu:'),
                'error': gettext('Hata:'),
                'unknown': gettext('Bilinmeyen hata'),
                'summary': gettext('Toplam {total} sonuç bulundu. İlk {shown} kayıt aşağıda gösteriliyor.'),
                'no_results': gettext('Bu arama için sonuç bulunamadı. Kriterleri genişletmeyi deneyin.'),
                'no_title': gettext('Başlık Yok'),
                'no_author': gettext('Yazar Yok'),
                'no_source': gettext('Kaynak Yok'),
                'no_abstract': gettext('Özet yok'),
                'sending': gettext('Gönderiliyor...'),
                'cancelling': gettext('İptal ediliyor...'),
                'previous': gettext('Önceki arama sonuçları'),
                'value_placeholder': gettext('Değer girin...'),
                'remove_rule': gettext('Kriteri kaldır'),
                'biblio_starting': gettext('Başlatılıyor...'),
                'biblio_started': gettext('Bibliometrik analiz kuyruğa alındı. Demo rapor (3 grafik PDF) hazır olduğunda e-postanıza gönderilecek.'),
                'biblio_exists': gettext('Bu arama için zaten bir analiz mevcut.'),
                'biblio_link': gettext('Bibliometrik analizlerim'),
            },
        },
    })


@login_required
@require_GET
def pubmed_job_status(request, job_id):
    """AJAX: job durumu + demo sonuçlar (ilk 3)."""
    job = get_object_or_404(PubMedSearchJob, id=job_id, user=request.user)
    response = {'status': job.status, 'total_results': job.total_results}
    if job.status == 'pending':
        from analizdestek.job_queue import get_queue_position
        response['queue_position'] = get_queue_position('pubmed', str(job.id))
    elif job.status == 'completed':
        response['demo_results'] = [
            {k: r.get(k, '') for k in ('title', 'authors', 'year', 'journal', 'type', 'abstract', 'pmid')}
            for r in (job.demo_results or [])[:3]
        ]
        response['has_demo'] = bool(job.demo_results)
    elif job.status == 'failed':
        response['error'] = job.error_message
    return JsonResponse(response)


@login_required
@require_GET
def pubmed_download_excel(request, job_id):
    """demo_results JSON'ından Excel (.xlsx) üretir."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    job = get_object_or_404(PubMedSearchJob, id=job_id, user=request.user)
    records = job.demo_results or []
    if not records:
        raise Http404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = gettext('PubMed Sonuçları')[:31]
    headers = [pgettext('sıra numarası', 'No'), gettext('Başlık'), gettext('Yazarlar'), gettext('Dergi'), gettext('Yıl'),
               'DOI', 'PMID', gettext('Tür'), gettext('Kurumlar'), gettext('Anahtar Kelimeler'),
               gettext('MeSH Terimleri'), gettext('Özet')]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = PatternFill(start_color='2D3748', end_color='2D3748', fill_type='solid')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)

    for i, r in enumerate(records, 1):
        ws.append([
            i, r.get('title', ''), r.get('authors', ''), r.get('journal', ''), r.get('year', ''),
            r.get('doi', ''), r.get('pmid', ''), r.get('type', ''), r.get('institutions', ''),
            ', '.join(r.get('keywords') or []), ', '.join(r.get('mesh_terms') or []), r.get('abstract', ''),
        ])
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{_safe_filename(job, "xlsx")}"'
    wb.save(response)
    return response


@login_required
@require_GET
def pubmed_download_txt(request, job_id):
    """demo_results'dan TXT dosyası üretir."""
    from .services.job_runner import _generate_pubmed_results_txt
    job = get_object_or_404(PubMedSearchJob, id=job_id, user=request.user)
    records = job.demo_results or []
    if not records:
        raise Http404
    response = HttpResponse(_generate_pubmed_results_txt(records, job), content_type='text/plain; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{_safe_filename(job, "txt")}"'
    return response


@login_required
@require_POST
def pubmed_cancel(request, job_id):
    """Devam eden taramayı iptal eder."""
    job = get_object_or_404(PubMedSearchJob, id=job_id, user=request.user)
    if job.status in ('pending', 'running'):
        job.status = 'failed'
        job.error_message = gettext('Kullanıcı tarafından iptal edildi.')
        job.save(update_fields=['status', 'error_message'])
    return JsonResponse({'success': True})


@login_required
@require_POST
def pubmed_send_demo_email(request, job_id):
    """Demo sonuçları kullanıcının kayıtlı e-postasına gönderir."""
    job = get_object_or_404(PubMedSearchJob, id=job_id, user=request.user)
    if job.status != 'completed' or not job.demo_results:
        return JsonResponse({'error': gettext('Sonuçlar henüz hazır değil.')}, status=400)
    if job.demo_email_sent:
        return JsonResponse({'error': gettext('Demo sonuçlar zaten gönderildi.')}, status=400)

    from .services.job_runner import send_demo_email_async
    send_demo_email_async(job.id)
    return JsonResponse({'success': True, 'message': gettext('Demo sonuçların {email} adresine gönderilmesi başlatıldı.').format(email=request.user.email)})
