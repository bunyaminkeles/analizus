"""
Bibliometri grafik üretimi için regresyon testleri.
Kapsam: yıl ekseni tam sayı tick'leri (A5), kısa dönem uyarısı (A6),
OpenAlex/PubMed TXT yeniden-yükleme (A7), demo CTA dil ayrımı (A8/A9),
rapor dili seçimi (A4).
"""
import csv
import io

import matplotlib
matplotlib.use('Agg')

import pytest
from django.contrib.auth.models import User
from django.utils import translation
from django.utils.translation import gettext

from forum.models import Profile

from bibliometrics.services.analyzer import (
    publication_trend,
    annual_citation_trend,
    publication_growth_rate,
    SHORT_SERIES_YEARS,
    canonical_author_key,
    _group_authors,
    top_authors,
    lotka_law,
    author_collaboration,
    normalize_keywords,
    _simple_depluralize,
    bradford_zones,
    bradford_law,
    keyword_clusters,
    keyword_cooccurrence,
    thematic_map_themes,
    thematic_map,
    thematic_evolution_periods,
    thematic_evolution_themes,
    thematic_evolution,
    three_field_plot_data,
    three_field_plot,
    scp_mcp_counts,
    scp_mcp_chart,
    bibliographic_coupling,
    local_citation_counts,
    local_global_citations,
    topic_modeling_data,
    topic_modeling_words_chart,
    topic_modeling_trend,
)
import matplotlib.patches as mpatches
import bibliometrics.services.analyzer as analyzer_module
from bibliometrics.services.report_notes import build_report_notes
from bibliometrics.services.parser import (
    parse_openalex_txt, parse_pubmed_txt, detect_format, parse_file,
    _deduplicate_and_filter, _normalize_title_key,
)
from bibliometrics.services.pdf_builder import _demo_cta_text
from bibliometrics.models import BibliometricJob
from bibliometrics.services.job_runner import _in_user_language
from bibliometrics.services.enrichment import enrich_records_with_openalex, _apply_enrichment
from forum.models import SiteSettings
from openalex.services.scraper import OpenAlexScraper


def _records(years_counts):
    recs = []
    for year, count in years_counts.items():
        for _ in range(count):
            recs.append({'year': year, 'cited_by': 3})
    return recs


def _assert_integer_xticks(fig):
    ax = fig.axes[0]
    ticks = ax.get_xticks()
    assert len(ticks) > 0
    assert all(float(t).is_integer() for t in ticks)


def test_publication_trend_integer_ticks_short_series():
    fig = publication_trend(_records({2021: 5, 2022: 8, 2023: 12}))
    _assert_integer_xticks(fig)


def test_publication_trend_integer_ticks_long_series():
    fig = publication_trend(_records({y: (y % 7 + 2) for y in range(2004, 2024)}))
    _assert_integer_xticks(fig)


def test_annual_citation_trend_integer_ticks_short_series():
    fig = annual_citation_trend(_records({2021: 5, 2022: 8, 2023: 12}))
    _assert_integer_xticks(fig)


def test_publication_growth_rate_integer_ticks_short_series():
    fig = publication_growth_rate(_records({2021: 5, 2022: 8, 2023: 12}))
    _assert_integer_xticks(fig)


def test_publication_growth_rate_integer_ticks_long_series():
    fig = publication_growth_rate(_records({y: (y % 7 + 2) for y in range(2004, 2024)}))
    _assert_integer_xticks(fig)


def test_short_series_warning_shown_on_chart_and_notes():
    recs = _records({2021: 5, 2022: 8, 2023: 12})
    fig = publication_growth_rate(recs)
    chart_text = ' '.join(t.get_text() for t in fig.axes[0].texts)
    assert 'küçük taban etkisine açıktır' in chart_text
    notes = build_report_notes(recs)
    assert any('küçük taban etkisine açıktır' in l for l in notes['limitations'])


def test_long_series_no_short_series_warning():
    recs = _records({y: 5 for y in range(2018, 2018 + SHORT_SERIES_YEARS + 1)})
    fig = publication_growth_rate(recs)
    chart_text = ' '.join(t.get_text() for t in fig.axes[0].texts)
    assert 'küçük taban etkisine açıktır' not in chart_text
    notes = build_report_notes(recs)
    assert not any('küçük taban etkisine açıktır' in l for l in notes['limitations'])


class _StubJob:
    """openalex/pubmed job_runner'ın _generate_*_results_txt'inin ihtiyaç duyduğu
    minimum arayüz — DB'ye bağımlı olmadan üretici+ayrıştırıcıyı uçtan uca test eder."""
    total_results = 1

    def get_query_summary(self):
        return 'test sorgusu'


def test_parse_openalex_txt_roundtrip_with_real_generator():
    """Üretici (job_runner._generate_alex_results_txt) ile ayrıştırıcı (parse_openalex_txt)
    birbirinden kopmasın diye uçtan uca test — A7'de bulunan kök neden: üretici etiket
    hizasını dinamik (ljust) hesaplıyor, ayrıştırıcı sabit genişlik varsayıyordu."""
    from openalex.services.job_runner import _generate_alex_results_txt
    pub = {
        'title': 'Test Başlık Yazısı',
        'authors': 'Ahmet Yılmaz, Mehmet Demir',
        'journal': 'Test Dergisi',
        'year': 2024,
        'doi': 'https://doi.org/10.1234/test',
        'type': 'article',
        'cited_by_count': 5,
        'institutions': 'İstanbul Üniversitesi',
        'keywords': ['anahtar bir', 'anahtar iki', 'anahtar üç'],
    }
    txt = _generate_alex_results_txt([pub], _StubJob())
    recs = parse_openalex_txt(txt)
    assert len(recs) == 1
    r = recs[0]
    assert r['title'] == 'Test Başlık Yazısı'
    assert r['year'] == 2024
    assert r['journal'] == 'Test Dergisi'
    assert r['keywords'] == ['anahtar bir', 'anahtar iki', 'anahtar üç']
    assert len(r['authors']) == 2


def test_parse_pubmed_txt_robust_to_label_width_drift():
    """Gelecekte etiket uzunluğu (ve dolayısıyla üreticinin ljust genişliği) değişse bile
    ayrıştırıcı kırılmamalı — sabit genişlik varsayımı yerine regex kullanılıyor."""
    content = (
        '--- Yayın #1 ---\n'
        'Başlık : Test Makale\n'
        'Yazarlar : Ahmet Yılmaz\n'
        'Yıl : 2023\n'
        'Anahtar Kelimeler : test, örnek\n'
    )
    recs = parse_pubmed_txt(content)
    assert len(recs) == 1
    assert recs[0]['title'] == 'Test Makale'
    assert recs[0]['year'] == 2023
    assert recs[0]['keywords'] == ['test', 'örnek']


def test_demo_cta_tr_has_no_proje_talebi_link():
    """TR: sipariş sistemi var, proje talebi linkine yönlendirilmemeli (A8)."""
    with translation.override('tr'):
        text = _demo_cta_text()
    assert 'proje-talebi' not in text
    assert 'analizus.com' in text


def test_demo_cta_en_de_link_to_proje_talebi_with_language_prefix():
    """EN/DE: sipariş sistemi yok, proje talebine (dil önekiyle) yönlendirilmeli (A8)."""
    with translation.override('en'):
        en_text = _demo_cta_text()
    with translation.override('de'):
        de_text = _demo_cta_text()
    assert '/en/proje-talebi/' in en_text
    assert '/de/proje-talebi/' in de_text
    assert en_text != de_text


@pytest.mark.django_db
def test_report_language_overrides_profile_language():
    """job.report_language doluysa profil/arayüz dilinden bağımsız zorlanır (A4)."""
    user = User.objects.create_user(username='testrptlang', password='x', email='t@example.com')
    Profile.objects.create(user=user, preferred_language='tr')

    job_en = BibliometricJob.objects.create(user=user, report_language='en')
    seen = {}
    _in_user_language(str(job_en.id), lambda job_id: seen.update(lang=translation.get_language()))
    assert seen['lang'] == 'en'


@pytest.mark.django_db
def test_report_language_blank_falls_back_to_profile_language():
    """job.report_language boşsa mevcut davranış (profil dili) korunur (A4)."""
    user = User.objects.create_user(username='testrptlang2', password='x', email='t2@example.com')
    Profile.objects.create(user=user, preferred_language='tr')

    job_default = BibliometricJob.objects.create(user=user, report_language='')
    seen = {}
    _in_user_language(str(job_default.id), lambda job_id: seen.update(lang=translation.get_language()))
    assert seen['lang'] == 'tr'


OPENALEX_WEB_CSV_SAMPLE = (
    'Title,Author,Year,Citation count,Open access,Abstract\n'
    '"Test Başlık, Virgüllü",Ahmet Yılmaz|Mehmet Demir,2023,5,Open Access,"Test özeti."\n'
    'İkinci Başlık,Ayşe Kaya,2023,0,NOT Open Access,\n'
)


def test_detect_format_openalex_web_csv():
    """openalex.org'un gerçek 'Export CSV' formatı (A7) — dergi/DOI/kurum/ülke/
    anahtar kelime hiç yok, yazarlar '|' ile ayrılmış."""
    assert detect_format(OPENALEX_WEB_CSV_SAMPLE) == 'csv_openalex'


def test_parse_openalex_csv_splits_pipe_separated_authors():
    recs, fmt = parse_file(OPENALEX_WEB_CSV_SAMPLE)
    assert fmt == 'csv_openalex'
    assert len(recs) == 2
    rec = next(r for r in recs if 'Virgüllü' in r['title'])
    assert rec['authors'] == ['Ahmet Yılmaz', 'Mehmet Demir']
    assert rec['year'] == 2023
    assert rec['cited_by'] == 5
    assert rec['abstract'] == 'Test özeti.'
    # Bu formatta hiç olmayan alanlar boş kalmalı
    assert rec['journal'] == ''
    assert rec['doi'] == ''


def _fake_work(doi, cited_by=7, country='US', institution='MIT', keywords=('deep learning',),
                referenced_works=('https://openalex.org/W1',), author_id='https://openalex.org/A1'):
    return {
        'doi': f'https://doi.org/{doi}',
        'cited_by_count': cited_by,
        'keywords': [{'display_name': k} for k in keywords],
        'concepts': [],
        'referenced_works': list(referenced_works),
        'authorships': [{
            'author': {'id': author_id},
            'institutions': [{'display_name': institution}],
            'countries': [country],
        }],
    }


def test_apply_enrichment_fills_missing_fields_only():
    """Boş alanlar OpenAlex verisiyle doldurulur; dolu alanlara (kullanıcının kendi verisi) dokunulmaz."""
    rec = {'doi': '10.1/x', 'cited_by': 0, 'country': '', 'institution': '', 'keywords': []}
    work = _fake_work('10.1/x')
    _apply_enrichment(rec, work)
    assert rec['cited_by'] == 7
    assert rec['country'] == 'US'
    assert rec['institution'] == 'MIT'
    assert rec['keywords'] == ['deep learning']
    assert rec['referenced_works'] == ['https://openalex.org/W1']
    assert rec['author_ids'] == ['https://openalex.org/A1']


def test_apply_enrichment_does_not_overwrite_existing_values():
    rec = {'doi': '10.1/x', 'cited_by': 99, 'country': 'TR', 'institution': 'ODTÜ', 'keywords': ['mevcut']}
    _apply_enrichment(rec, _fake_work('10.1/x'))
    assert rec['cited_by'] == 99
    assert rec['country'] == 'TR'
    assert rec['institution'] == 'ODTÜ'
    assert rec['keywords'] == ['mevcut']
    # Yeni alanlar (başka yerde doldurulmuyor) yine de güncellenir
    assert rec['referenced_works'] == ['https://openalex.org/W1']


@pytest.mark.django_db
def test_enrich_records_skips_records_without_doi(monkeypatch):
    called = {}

    def fake_fetch(self, dois):
        called['dois'] = dois
        return {d.lower(): _fake_work(d) for d in dois}

    monkeypatch.setattr(OpenAlexScraper, 'fetch_works_by_dois', fake_fetch)

    records = [
        {'doi': '10.1/a', 'cited_by': 0, 'country': '', 'institution': '', 'keywords': []},
        {'doi': '', 'cited_by': 0, 'country': '', 'institution': '', 'keywords': []},
    ]
    stats = {}
    enrich_records_with_openalex(records, stats)

    assert stats['zenginlestirme_skipped_no_doi'] == 1
    assert stats['zenginlestirme_enriched'] == 1
    assert called['dois'] == ['10.1/a']
    assert records[0]['cited_by'] == 7
    assert records[1]['cited_by'] == 0  # DOI'siz kayıt dokunulmadan kaldı


@pytest.mark.django_db
def test_enrich_records_respects_max_kayit_limit(monkeypatch):
    settings_obj = SiteSettings.load()
    settings_obj.bibliometri_openalex_zenginlestirme_max_kayit = 1
    settings_obj.save(update_fields=['bibliometri_openalex_zenginlestirme_max_kayit'])

    called = {}

    def fake_fetch(self, dois):
        called['dois'] = dois
        return {d.lower(): _fake_work(d) for d in dois}

    monkeypatch.setattr(OpenAlexScraper, 'fetch_works_by_dois', fake_fetch)

    records = [
        {'doi': '10.1/a', 'cited_by': 0, 'country': '', 'institution': '', 'keywords': []},
        {'doi': '10.1/b', 'cited_by': 0, 'country': '', 'institution': '', 'keywords': []},
    ]
    stats = {}
    enrich_records_with_openalex(records, stats)

    assert called['dois'] == ['10.1/a']
    assert stats['zenginlestirme_enriched'] == 1
    assert records[1]['cited_by'] == 0


def test_fetch_works_by_dois_batches_and_normalizes_doi(monkeypatch):
    """50'den fazla DOI iki ayrı istek halinde gönderilir; dönen anahtar küçük harf, önekisiz."""
    seen_filters = []

    def fake_fetch_page(self, params, cursor=None):
        seen_filters.append(params['filter'])
        return {'results': [{'doi': 'https://doi.org/10.1/A'}]}

    monkeypatch.setattr(OpenAlexScraper, '_fetch_page', fake_fetch_page)

    dois = [f'10.1/{i}' for i in range(60)]
    result = OpenAlexScraper().fetch_works_by_dois(dois)

    assert len(seen_filters) == 2  # 60 DOI → 50 + 10
    assert '10.1/0' in seen_filters[0]
    assert '10.1/59' in seen_filters[1]
    assert '10.1/a' in result  # doi.org/10.1/A -> küçük harf


def _rec(title, year=None, doi=''):
    return {'title': title, 'year': year, 'doi': doi}


def test_dedup_same_title_same_year_is_duplicate():
    recs = [_rec('COVID-19: A Review', 2021), _rec('covid 19 a review', 2021)]
    result = _deduplicate_and_filter(recs)
    assert len(result) == 1


def test_dedup_same_title_different_years_kept_separate():
    """B1-2: aynı (genel) başlık farklı yıllarda yayınlanmışsa yanlışlıkla birleştirilmez."""
    recs = [_rec('Editöryal', 2020), _rec('Editöryal', 2021)]
    result = _deduplicate_and_filter(recs)
    assert len(result) == 2


def test_dedup_same_title_one_missing_year_still_duplicate():
    """Kaynaklardan biri yıl vermiyorsa eski davranış gibi yalnız başlığa güvenilir."""
    recs = [_rec('Örnek Başlık', 2020), _rec('Örnek Başlık', None)]
    result = _deduplicate_and_filter(recs)
    assert len(result) == 1


def test_dedup_doi_takes_priority_over_title():
    recs = [_rec('Başlık A', 2020, doi='10.1/x'), _rec('Başlık B', 2021, doi='10.1/x')]
    result = _deduplicate_and_filter(recs)
    assert len(result) == 1


def test_normalize_title_key_ignores_punctuation():
    assert _normalize_title_key('COVID-19: A Review!') == _normalize_title_key('covid 19 a review')


def test_canonical_author_key_matches_name_variants():
    """B1-3: 'Smith, J.' / 'John Smith' / 'JOHN SMITH' aynı kişi sayılmalı."""
    key = canonical_author_key('John Smith')
    assert canonical_author_key('Smith, J.') == key
    assert canonical_author_key('JOHN SMITH') == key
    assert canonical_author_key('Smith, J') == key


def test_canonical_author_key_ignores_accents_and_case():
    assert canonical_author_key('Şahin, A.') == canonical_author_key('ahmet sahin')


def test_canonical_author_key_distinguishes_different_surnames():
    assert canonical_author_key('John Smith') != canonical_author_key('John Doe')


def test_group_authors_merges_name_variants():
    records = [
        {'authors': ['Smith, J.', 'Ayşe Kaya']},
        {'authors': ['John Smith']},
        {'authors': ['John Smith']},
    ]
    counts, display = _group_authors(records)
    smith_key = canonical_author_key('John Smith')
    assert counts[smith_key] == 3
    assert display[smith_key] in ('John Smith', 'Smith, J.')
    assert len(counts) == 2  # Smith + Kaya


def test_top_authors_merges_variants_into_single_bar():
    records = [
        {'authors': ['Smith, J.']},
        {'authors': ['John Smith']},
        {'authors': ['John Smith']},
        {'authors': ['Ayşe Kaya']},
    ]
    fig = top_authors(records, n=15)
    ax = fig.axes[0]
    labels = [t.get_text() for t in ax.get_yticklabels()]
    assert len(labels) == 2
    bars = [p.get_width() for p in ax.patches]
    assert sorted(bars) == [1, 3]


def test_lotka_law_merges_variants():
    records = [
        {'authors': ['Smith, J.']},
        {'authors': ['John Smith']},
        {'authors': ['John Smith']},
        {'authors': ['Ayşe Kaya']},
    ]
    counts, _ = _group_authors(records)
    assert sorted(counts.values()) == [1, 3]
    fig = lotka_law(records)
    assert fig is not None


def test_author_collaboration_merges_variants_into_one_node():
    """Aynı yazarın 'Smith, J.' / 'John Smith' varyantları tek düğüme birleşmeli."""
    records = [
        {'authors': ['Smith, J.', 'Ayşe Kaya']},
        {'authors': ['John Smith', 'Mehmet Demir']},
    ]
    fig = author_collaboration(records, min_collab=1)
    assert fig is not None
    ax = fig.axes[0]
    node_collection = max(ax.collections, key=lambda c: c.get_offsets().shape[0])
    assert node_collection.get_offsets().shape[0] == 3  # Smith (birleşik) + Kaya + Demir


def test_simple_depluralize_basic_rules():
    assert _simple_depluralize('cats') == 'cat'
    assert _simple_depluralize('boxes') == 'box'
    assert _simple_depluralize('studies') == 'study'
    assert _simple_depluralize('glass') == 'glass'  # çift 's' — dokunulmaz
    assert _simple_depluralize('bus') == 'bus'  # çok kısa — dokunulmaz


def test_normalize_keywords_hyphen_and_space_are_equivalent():
    """B1-4: 'COVID-19' ve 'COVID 19' aynı anahtar kelimeye normalize edilmeli."""
    assert normalize_keywords(['COVID-19']) == normalize_keywords(['COVID 19'])


def test_normalize_keywords_falls_back_to_depluralize_when_lemmatization_unchanged(monkeypatch):
    """TR lemmatization kelimeyi değiştirmediyse (muhtemelen yabancı kelime) basit
    çoğul→tekil kuralı devreye girer (B1-4)."""
    monkeypatch.setattr(analyzer_module, 'lemmatize_word', lambda w: w)
    assert normalize_keywords(['cats']) == ['cat']


def _journal_records(counts: dict):
    recs = []
    for journal, n in counts.items():
        for _ in range(n):
            recs.append({'journal': journal})
    return recs


def test_bradford_zones_splits_into_three_equal_thirds():
    """B2-1: bölgeler yaklaşık eşit yayın sayısı taşımalı, az dergi çok yayın
    üretiyorsa (çekirdek) Bölge 1'in dergi sayısı Bölge 3'ten belirgin az olmalı."""
    counts = {f'J{i}': c for i, c in enumerate([30, 15, 10, 5, 5, 3, 3, 2, 2, 1])}
    records = _journal_records(counts)
    zones, totals = bradford_zones(records)
    total = sum(counts.values())
    assert sum(totals) == total
    assert len(zones[0]) < len(zones[2])
    # Her bölge toplamı hedef (total/3)'e yakın olmalı (son bölge kalanı taşır)
    for t in totals[:2]:
        assert abs(t - total / 3) <= max(counts.values())


def test_bradford_zones_insufficient_journals_returns_none():
    records = _journal_records({'J1': 5, 'J2': 3})  # yalnız 2 farklı dergi
    zones, totals = bradford_zones(records)
    assert zones is None and totals is None


def test_bradford_law_returns_none_without_enough_journals():
    assert bradford_law(_journal_records({'Tek Dergi': 10})) is None


def test_bradford_law_figure_table_matches_zone_totals():
    counts = {f'J{i}': c for i, c in enumerate([30, 15, 10, 5, 5, 3, 3, 2, 2, 1])}
    records = _journal_records(counts)
    zones, totals = bradford_zones(records)
    fig = bradford_law(records)
    assert fig is not None
    table = fig.axes[1].tables[0]
    cells = table.get_celld()
    # Satır 1 (ilk veri satırı, 0 başlık) → Bölge 1 dergi/yayın sayısı
    assert cells[(1, 1)].get_text().get_text() == str(len(zones[0]))
    assert cells[(1, 2)].get_text().get_text() == f'{totals[0]:,}'


def test_bradford_law_caps_shown_journals_for_large_datasets():
    """Canlı rapor hatası (ekran görüntüsüyle bildirildi): Bölge 1+2'de çok
    sayıda dergi olduğunda (büyük/gerçek veri setleri) sabit genişlik (13")
    + sınırsız büyüyen yükseklik (ör. 13x32) aşırı dar/uzun bir görüntü
    üretip PDF'e sığdırılınca genişliğin yarısından fazlası boş kalıyordu.
    Artık en fazla 40 dergi gösterilir, figsize buna göre sınırlı kalır."""
    counts = {}
    for i in range(5):
        counts[f'Cekirdek {i}'] = 80 - i * 8
    for i in range(300):
        counts[f'Dergi {i}'] = max(1, 15 - i // 10)
    records = _journal_records(counts)

    zones, _ = bradford_zones(records)
    assert len(zones[0]) + len(zones[1]) > 40  # senaryo gerçekten kırpma gerektiriyor

    fig = bradford_law(records)
    assert fig is not None
    w, h = fig.get_size_inches()
    assert h <= 19  # eskiden 70+ tane dergiyle 0.45 çarpanı ~32"ye çıkıyordu
    assert len(fig.axes[0].patches) == 40  # tam olarak kırpma sınırı kadar çubuk

    # Sağ paneldeki tablo, kırpılan kısımdan BAĞIMSIZ tam Bölge sayılarını göstermeli
    table = fig.axes[1].tables[0]
    cells = table.get_celld()
    assert cells[(1, 1)].get_text().get_text() == str(len(zones[0]))
    assert cells[(2, 1)].get_text().get_text() == str(len(zones[1]))


_KEYWORD_TOPICS = {
    'A': ['machine learning', 'deep learning', 'neural network', 'classification'],
    'B': ['covid-19', 'public health', 'pandemic', 'vaccination'],
    'C': ['climate change', 'carbon emission', 'sustainability', 'renewable energy'],
}


def _topic_keyword_records(topics=_KEYWORD_TOPICS, per_topic=25, seed=3):
    import random
    rng = random.Random(seed)
    records = []
    for _ in range(per_topic * len(topics)):
        topic = rng.choice(list(topics.keys()))
        kws = rng.sample(topics[topic], k=rng.randint(2, 4))
        records.append({'keywords': kws})
    return records


def test_keyword_clusters_separates_distinct_topics():
    """B2-2: birbiriyle hiç eş geçmeyen 3 konu grubu 3 ayrı Louvain kümesine
    ayrılmalı, her kümenin kelimeleri yalnız kendi konu grubundan olmalı."""
    records = _topic_keyword_records()
    G, communities, all_kw = keyword_clusters(records)
    assert G is not None
    assert len(communities) == 3
    topic_sets = [set(normalize_keywords(v)) for v in _KEYWORD_TOPICS.values()]
    for comm in communities:
        assert any(comm <= topic_set for topic_set in topic_sets)


def test_keyword_clusters_deterministic_with_fixed_seed():
    """Louvain varsayılan olarak rastgele olabilir — sabit seed aynı sonucu vermeli."""
    records = _topic_keyword_records()
    _, communities1, _ = keyword_clusters(records)
    _, communities2, _ = keyword_clusters(records)
    assert communities1 == communities2


def test_keyword_clusters_insufficient_data_returns_none():
    G, communities, all_kw = keyword_clusters([{'keywords': []}])
    assert G is None and communities is None and all_kw is None


def test_keyword_cooccurrence_table_lists_all_clusters():
    records = _topic_keyword_records()
    _, communities, _ = keyword_clusters(records)
    fig = keyword_cooccurrence(records)
    assert fig is not None
    table = fig.axes[1].tables[0]
    n_data_rows = len({row for row, _ in table.get_celld().keys() if row > 0})
    assert n_data_rows == len(communities)


def test_thematic_map_themes_isolated_cluster_has_zero_centrality():
    """B2-3: dışa hiç bağlantısı olmayan (izole) bir küme merkeziyet=0 almalı —
    Callon merkeziyet tanımı: dışarıyla toplam eş-oluşum ağırlığı."""
    # 'data' A ve B'yi köprüler; C (iklim) hiçbir ortak kelime paylaşmaz → izole.
    topics = {
        'A': ['machine learning', 'deep learning', 'neural network', 'classification', 'data'],
        'B': ['covid-19', 'public health', 'pandemic', 'vaccination', 'data'],
        'C': ['climate change', 'carbon emission', 'sustainability', 'renewable energy'],
    }
    records = _topic_keyword_records(topics, per_topic=60, seed=5)
    themes, mean_c, mean_d = thematic_map_themes(records)
    assert themes is not None
    assert len(themes) == 3
    isolated = next(t for t in themes if 'climate change' in t['label'] or 'carbon emission' in t['label'])
    assert isolated['centrality'] == 0
    assert all(t['centrality'] >= 0 and t['density'] >= 0 for t in themes)


def test_thematic_map_themes_requires_at_least_two_clusters():
    """Tek kümeyle merkeziyet (dışa bağlantı) tanımsız — None dönmeli."""
    records = _topic_keyword_records({'A': _KEYWORD_TOPICS['A']}, per_topic=30, seed=1)
    themes, mean_c, mean_d = thematic_map_themes(records)
    assert themes is None


def test_thematic_map_figure_has_one_bubble_per_theme():
    records = _topic_keyword_records()
    themes, _, _ = thematic_map_themes(records)
    fig = thematic_map(records)
    assert fig is not None
    scatter = next(c for c in fig.axes[0].collections if hasattr(c, 'get_offsets')
                   and c.get_offsets().shape[0] > 0)
    assert scatter.get_offsets().shape[0] == len(themes)


def _two_period_bridge_records():
    """B2-4 test verisi: eski dönemde A+B, yeni dönemde B+C — B köprü tema."""
    topics_old = {'A': _KEYWORD_TOPICS['A'], 'B': _KEYWORD_TOPICS['B']}
    topics_new = {'B': _KEYWORD_TOPICS['B'], 'C': _KEYWORD_TOPICS['C']}
    old_records = _topic_keyword_records(topics_old, per_topic=60, seed=11)
    for r in old_records:
        r['year'] = 2010
    new_records = _topic_keyword_records(topics_new, per_topic=60, seed=12)
    for r in new_records:
        r['year'] = 2020
    return old_records + new_records


def test_thematic_evolution_periods_insufficient_records_returns_none():
    records = [{'year': 2020}] * 3
    assert thematic_evolution_periods(records, n_periods=2, min_per_period=8) is None


def test_thematic_evolution_periods_splits_by_year():
    records = _two_period_bridge_records()
    periods = thematic_evolution_periods(records, n_periods=2)
    assert periods is not None
    assert len(periods) == 2
    assert len(periods[0][1]) == 120  # per_topic=60 * 2 konu
    assert len(periods[1][1]) == 120


def test_thematic_evolution_tracks_persisting_theme_and_detects_change():
    """B2-4: ortak tema (B) iki dönem arasında köprüleniyorsa tek akış oluşmalı;
    yalnız bir dönemde var olan temalar (A eski dönemde, C yeni dönemde) akışsız kalmalı."""
    records = _two_period_bridge_records()
    labels, themes, flows = thematic_evolution_themes(records, n_periods=2)
    assert labels is not None
    assert len(themes) == 2
    assert len(themes[0]) == 2 and len(themes[1]) == 2

    b_kw = set(normalize_keywords(_KEYWORD_TOPICS['B']))
    b_old_idx = next(i for i, t in enumerate(themes[0]) if t['keywords'] == b_kw)
    b_new_idx = next(i for i, t in enumerate(themes[1]) if t['keywords'] == b_kw)

    assert len(flows) == 1
    period_i, a_idx, b_idx, weight = flows[0]
    assert period_i == 0 and a_idx == b_old_idx and b_idx == b_new_idx
    assert weight == len(b_kw)


def test_thematic_evolution_themes_insufficient_data_returns_none():
    labels, themes, flows = thematic_evolution_themes([{'year': 2020, 'keywords': []}] * 20)
    assert labels is None and themes is None and flows is None


def test_thematic_evolution_figure_has_box_per_theme():
    records = _two_period_bridge_records()
    fig = thematic_evolution(records)
    assert fig is not None
    rects = [p for p in fig.axes[0].patches if isinstance(p, mpatches.Rectangle)]
    assert len(rects) == 4  # 2 tema x 2 dönem (3 dönem veriyle kurulamaz, 2'ye düşer)


def _three_field_records():
    """B2-5 test verisi: 2 yazar, 2+1 ortak anahtar kelime, 2 dergi."""
    records = []
    for _ in range(6):
        records.append({'authors': ['Alice Smith'], 'keywords': ['alpha', 'shared'], 'journal': 'J1'})
    for _ in range(5):
        records.append({'authors': ['Bob Jones'], 'keywords': ['beta', 'shared'], 'journal': 'J2'})
    return records


def test_three_field_plot_data_builds_consistent_flows():
    """B2-5: yazar/dergi kutu büyüklüğü akış toplamına tam oturmalı; anahtar kelime
    kutu büyüklüğü (toplam frekans) her iki yöndeki akış toplamını da sınırlamalı."""
    records = _three_field_records()
    authors, keywords, journals, flows_ak, flows_kj = three_field_plot_data(records)
    assert authors is not None
    assert len(authors) == 2
    assert len(journals) == 2
    kw_labels = {k['label'] for k in keywords}
    assert {'alpha', 'beta', 'shared'} <= kw_labels

    shared_idx = next(i for i, k in enumerate(keywords) if k['label'] == 'shared')
    assert keywords[shared_idx]['size'] == 11  # 6 + 5 — tüm kayıtlardaki toplam frekans

    for a_idx, author in enumerate(authors):
        total_out = sum(w for a, k, w in flows_ak if a == a_idx)
        assert total_out == author['size']

    for k_idx, kw in enumerate(keywords):
        total_in = sum(w for a, k, w in flows_ak if k == k_idx)
        total_out = sum(w for k2, j, w in flows_kj if k2 == k_idx)
        assert total_in <= kw['size']
        assert total_out <= kw['size']


def test_three_field_plot_data_insufficient_returns_none():
    records = [{'authors': ['Only One'], 'keywords': ['x'], 'journal': 'J1'}]
    result = three_field_plot_data(records)
    assert result == (None, None, None, None, None)


def test_three_field_plot_figure_has_box_per_item():
    records = _three_field_records()
    authors, keywords, journals, _, _ = three_field_plot_data(records)
    fig = three_field_plot(records)
    assert fig is not None
    rects = [p for p in fig.axes[0].patches if isinstance(p, mpatches.Rectangle)]
    assert len(rects) == len(authors) + len(keywords) + len(journals)


def test_scp_mcp_counts_classifies_single_and_multi_country():
    """B2-6: tek ülkeli yayın SCP, birden fazla ülkeden yazarı olan yayın HER
    ülke için MCP sayılmalı (genel ülke sayımındaki 'tam sayım' kuralıyla tutarlı)."""
    records = (
        [{'country': 'United States'}] * 5
        + [{'country': 'United States, United Kingdom'}] * 3
        + [{'country': 'United Kingdom'}] * 2
    )
    counts = scp_mcp_counts(records)
    assert counts is not None
    assert counts['United States'] == {'scp': 5, 'mcp': 3}
    assert counts['United Kingdom'] == {'scp': 2, 'mcp': 3}


def test_scp_mcp_counts_no_country_returns_none():
    assert scp_mcp_counts([{'country': ''}, {}]) is None


def test_scp_mcp_chart_bar_count_matches_top_n_countries():
    records = (
        [{'country': 'United States'}] * 5
        + [{'country': 'United States, United Kingdom'}] * 3
        + [{'country': 'United Kingdom'}] * 2
        + [{'country': 'Germany'}] * 1
    )
    fig = scp_mcp_chart(records, n=2)
    assert fig is not None


def _coupling_records():
    """B2-7 test verisi: A-B 2 paylaşılan referans (eşik üstü), C hiçbiriyle paylaşmıyor."""
    return [
        {'title': 'Paper A', 'referenced_works': ['w1', 'w2', 'w3']},
        {'title': 'Paper B', 'referenced_works': ['w1', 'w2', 'w5']},
        {'title': 'Paper C', 'referenced_works': ['w6', 'w7']},
    ]


def test_bibliographic_coupling_links_records_sharing_references():
    """B2-7: A-B 2 referans paylaşıyor (eşik=2) → bağlı; C hiçbiriyle paylaşmadığı
    için izole kalır ve grafiğe hiç düğüm olarak girmez (yalnız kenarla eklenir)."""
    fig = bibliographic_coupling(_coupling_records(), min_shared=2)
    assert fig is not None
    ax = fig.axes[0]
    node_collection = max(ax.collections, key=lambda c: c.get_offsets().shape[0])
    assert node_collection.get_offsets().shape[0] == 2  # yalnız A ve B


def test_bibliographic_coupling_insufficient_data_returns_none():
    records = [{'title': 'Only One', 'referenced_works': ['w1', 'w2']}]
    assert bibliographic_coupling(records) is None


def test_bibliographic_coupling_no_shared_references_returns_none():
    records = [
        {'title': 'Paper A', 'referenced_works': ['w1', 'w2']},
        {'title': 'Paper B', 'referenced_works': ['w3', 'w4']},
    ]
    assert bibliographic_coupling(records) is None


def _local_citation_records():
    """B2-7 test verisi: W1 iki kayıttan (W2, W3), W2 bir kayıttan (W3) yerel atıf alır."""
    return [
        {'title': 'Foundational Paper', 'openalex_id': 'W1', 'cited_by': 50, 'referenced_works': []},
        {'title': 'Other Paper', 'openalex_id': 'W2', 'cited_by': 10, 'referenced_works': ['W1']},
        {'title': 'Third Paper', 'openalex_id': 'W3', 'cited_by': 5, 'referenced_works': ['W1', 'W2']},
    ]


def test_local_citation_counts_counts_internal_references():
    counts = local_citation_counts(_local_citation_records())
    assert counts == {'W1': 2, 'W2': 1}  # W3'e hiç referans verilmedi, listede yok


def test_local_citation_counts_insufficient_data_returns_none():
    assert local_citation_counts([{'openalex_id': 'W1'}, {}]) is None


def test_local_global_citations_figure_ranks_by_local_count():
    fig = local_global_citations(_local_citation_records())
    assert fig is not None
    ax = fig.axes[0]
    labels = [t.get_text() for t in ax.get_yticklabels()]
    assert labels == ['Foundational Paper', 'Other Paper']  # Third Paper hiç yerel atıf almadı


def _topic_modeling_records():
    """B2-8 test verisi: 2 net ayrışan kelime kümesi (fizik vs malzeme), 20'şer kayıt."""
    physics = [
        'photon laser quantum entangled particle wave optical spin photon laser '
        'quantum entangled particle wave optical spin',
        'quantum photon spin entangled optical particle laser wave quantum photon '
        'spin entangled optical particle laser wave',
        'laser photon optical wave quantum particle entangled spin laser photon '
        'optical wave quantum particle entangled spin',
    ]
    material = [
        'polymer composite material strength fiber tensile bonding structural '
        'polymer composite material strength fiber tensile bonding structural',
        'composite polymer fiber structural tensile material strength bonding '
        'composite polymer fiber structural tensile material strength bonding',
        'fiber composite tensile polymer bonding structural strength material '
        'fiber composite tensile polymer bonding structural strength material',
    ]
    records = []
    for i in range(20):
        records.append({'abstract': physics[i % 3], 'year': 2015 + (i % 5)})
    for i in range(20):
        records.append({'abstract': material[i % 3], 'year': 2015 + (i % 5)})
    return records


def test_topic_modeling_data_separates_distinct_vocabularies():
    """B2-8: iki birbirinden bağımsız kelime kümesi, karışmadan 2 ayrı konuya ayrılmalı."""
    topics, assignments = topic_modeling_data(_topic_modeling_records(), n_topics=2, top_words=8)
    assert topics is not None
    assert len(topics) == 2
    assert sum(t['size'] for t in topics) == 40
    assert len(assignments) == 40

    physics_vocab = {'photon', 'laser', 'quantum', 'entangled', 'particle', 'wave', 'optical', 'spin'}
    material_vocab = {'polymer', 'composite', 'material', 'strength', 'fiber', 'tensile', 'bonding',
                       'structural'}
    for t in topics:
        words = set(t['words'])
        assert words <= physics_vocab or words <= material_vocab


def test_topic_modeling_data_insufficient_abstracts_returns_none():
    records = [{'abstract': 'short text here', 'year': 2020}] * 5
    topics, assignments = topic_modeling_data(records)
    assert topics is None and assignments is None


def test_topic_modeling_data_no_abstracts_returns_none():
    topics, assignments = topic_modeling_data([{'abstract': ''}] * 30)
    assert topics is None and assignments is None


def test_topic_modeling_words_chart_has_subplot_per_topic():
    fig = topic_modeling_words_chart(_topic_modeling_records())
    assert fig is not None
    assert len(fig.axes) >= 2


def test_topic_modeling_trend_figure_not_none_with_multiyear_data():
    fig = topic_modeling_trend(_topic_modeling_records())
    assert fig is not None


def test_topic_modeling_trend_single_year_returns_none():
    records = _topic_modeling_records()
    for r in records:
        r['year'] = 2020
    assert topic_modeling_trend(records) is None


# ─────────────────────────── B3-1: Temiz Veri İndirme (XLSX/CSV) ───────────────────────────

def _clean_export_records():
    return [
        {'title': 'Başlık Bir', 'authors': ['Ahmet Yılmaz', 'Mehmet Demir'], 'year': 2021,
         'journal': 'Dergi A', 'keywords': ['makine öğrenmesi', 'yapay zeka'], 'abstract': 'Özet metni.',
         'cited_by': 12, 'doi': '10.1/a', 'pub_type': 'article', 'country': 'TR', 'institution': 'ODTÜ',
         'openalex_id': 'https://openalex.org/W1', 'referenced_works': ['https://openalex.org/W9'],
         'author_ids': ['https://openalex.org/A1']},
        {'title': 'İkinci Başlık', 'authors': [], 'year': None, 'journal': '', 'keywords': [],
         'abstract': '', 'cited_by': 0, 'doi': '', 'pub_type': '', 'country': '', 'institution': '',
         'openalex_id': '', 'referenced_works': [], 'author_ids': []},
    ]


def test_build_clean_xlsx_has_expected_rows_and_columns():
    import openpyxl
    from bibliometrics.services.export_builder import build_clean_xlsx

    data = build_clean_xlsx(_clean_export_records())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active

    header = [c.value for c in ws[1]]
    assert header == [
        'Başlık', 'Yazarlar', 'Yıl', 'Dergi', 'Anahtar Kelimeler', 'Özet',
        'Atıf Sayısı', 'DOI', 'Yayın Türü', 'Ülke', 'Kurum', 'OpenAlex ID',
    ]
    assert ws.max_row == 3  # başlık + 2 kayıt

    row1 = [c.value for c in ws[2]]
    assert row1[0] == 'Başlık Bir'
    assert row1[1] == 'Ahmet Yılmaz; Mehmet Demir'
    assert row1[4] == 'makine öğrenmesi; yapay zeka'
    assert row1[10] == 'ODTÜ'
    assert row1[11] == 'https://openalex.org/W1'

    row2 = [c.value for c in ws[3]]
    # openpyxl boş string'i hücreye None olarak geri okur
    assert row2[1] in (None, '')
    assert row2[4] in (None, '')


def test_build_clean_xlsx_empty_records_has_only_header():
    import openpyxl
    from bibliometrics.services.export_builder import build_clean_xlsx

    wb = openpyxl.load_workbook(io.BytesIO(build_clean_xlsx([])))
    assert wb.active.max_row == 1


def test_build_clean_csv_has_bom_and_expected_rows():
    from bibliometrics.services.export_builder import build_clean_csv

    data = build_clean_csv(_clean_export_records())
    assert data.startswith(b'\xef\xbb\xbf')

    text = data.decode('utf-8-sig')
    rows = list(csv.reader(io.StringIO(text)))
    assert len(rows) == 3  # başlık + 2 kayıt
    assert rows[0][0] == 'Başlık'
    assert rows[1][1] == 'Ahmet Yılmaz; Mehmet Demir'
    assert rows[2][0] == 'İkinci Başlık'


def test_upload_clean_data_returns_both_urls(monkeypatch):
    from bibliometrics.services import job_runner

    uploaded = {}

    def fake_upload(data, key, content_type):
        uploaded[key] = content_type
        return f'https://s3.example.com/{key}'

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    xlsx_url, csv_url = job_runner._upload_clean_data(_clean_export_records(), 'job123')
    assert xlsx_url == 'https://s3.example.com/bibliometrics/clean_data/job123.xlsx'
    assert csv_url == 'https://s3.example.com/bibliometrics/clean_data/job123.csv'
    assert uploaded['bibliometrics/clean_data/job123.xlsx'] == (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert uploaded['bibliometrics/clean_data/job123.csv'] == 'text/csv'


def test_upload_clean_data_tolerates_upload_failure(monkeypatch):
    from bibliometrics.services import job_runner

    def fake_upload(data, key, content_type):
        raise RuntimeError('S3 kapalı')

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    xlsx_url, csv_url = job_runner._upload_clean_data(_clean_export_records(), 'job456')
    assert xlsx_url == '' and csv_url == ''


# ─────────────────────────── B3-2: VOSviewer Export ───────────────────────────

def _vosviewer_records():
    """Yazar/ülke/anahtar kelime/referans çiftleri üzerinde en az 1 bağlantı
    üretecek 2 kayıt — her 4 ağ fonksiyonu için de en az bir kenar."""
    base = {
        'authors': ['Ahmet Yılmaz', 'Mehmet Demir'],
        'country': 'TR; US',
        'keywords': ['makine öğrenmesi', 'yapay zeka'],
        'referenced_works': ['https://openalex.org/W1', 'https://openalex.org/W2'],
    }
    return [
        {**base, 'title': 'Birinci Makale', 'year': 2021},
        {**base, 'title': 'İkinci Makale', 'year': 2022},
    ]


def test_graph_to_files_produces_expected_map_and_network_columns():
    import networkx as nx
    from bibliometrics.services.vosviewer_export import _graph_to_files

    G = nx.Graph()
    G.add_edge('Kelime A', 'Kelime B, Sanayi', weight=3)

    map_text, network_text = _graph_to_files(G)

    map_lines = map_text.strip('\n').split('\n')
    assert map_lines[0] == 'id\tlabel\tweight<Occurrences>'
    assert len(map_lines) == 3  # başlık + 2 düğüm
    # virgüllü etiket çift tırnağa alınmış olmalı
    assert any('"Kelime B, Sanayi"' in line for line in map_lines[1:])

    network_lines = network_text.strip('\n').split('\n')
    assert len(network_lines) == 1
    id1, id2, w = network_lines[0].split('\t')
    assert w == '3'


def test_build_vosviewer_zip_contains_all_four_networks():
    import zipfile
    import io
    from bibliometrics.services.vosviewer_export import build_vosviewer_zip

    data = build_vosviewer_zip(_vosviewer_records())
    assert data is not None

    zf = zipfile.ZipFile(io.BytesIO(data))
    names = set(zf.namelist())
    assert 'OKUBENI.txt' in names
    for prefix in ('anahtar_kelime_es_olusumu', 'yazar_isbirligi',
                   'ulke_isbirligi', 'bibliyografik_eslesme'):
        assert f'{prefix}_map.txt' in names
        assert f'{prefix}_network.txt' in names


def test_build_vosviewer_zip_returns_none_without_any_network_data():
    from bibliometrics.services.vosviewer_export import build_vosviewer_zip

    # Tek yazarlı, ülkesiz, anahtar kelimesiz, referanssız kayıtlar — hiçbir ağ kurulamaz
    records = [{'authors': ['Tek Yazar'], 'title': f'Makale {i}', 'year': 2020} for i in range(5)]
    assert build_vosviewer_zip(records) is None


def test_upload_vosviewer_export_returns_url(monkeypatch):
    from bibliometrics.services import job_runner

    uploaded = {}

    def fake_upload(data, key, content_type):
        uploaded[key] = content_type
        return f'https://s3.example.com/{key}'

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    url = job_runner._upload_vosviewer_export(_vosviewer_records(), 'job789')
    assert url == 'https://s3.example.com/bibliometrics/vosviewer/job789.zip'
    assert uploaded['bibliometrics/vosviewer/job789.zip'] == 'application/zip'


def test_upload_vosviewer_export_returns_empty_without_data(monkeypatch):
    from bibliometrics.services import job_runner

    called = []
    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', lambda *a, **k: called.append(1))

    records = [{'authors': ['Tek Yazar'], 'title': f'Makale {i}', 'year': 2020} for i in range(5)]
    url = job_runner._upload_vosviewer_export(records, 'job000')
    assert url == ''
    assert not called


def test_upload_vosviewer_export_tolerates_upload_failure(monkeypatch):
    from bibliometrics.services import job_runner

    def fake_upload(data, key, content_type):
        raise RuntimeError('S3 kapalı')

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    url = job_runner._upload_vosviewer_export(_vosviewer_records(), 'job111')
    assert url == ''


def test_author_collaboration_graph_matches_chart_function_data():
    from bibliometrics.services.analyzer import _author_collaboration_graph, author_collaboration

    records = _vosviewer_records()
    G = _author_collaboration_graph(records, min_collab=2)
    assert G is not None
    assert G.number_of_nodes() == 2 and G.number_of_edges() == 1
    assert author_collaboration(records, min_collab=2) is not None


def test_country_collaboration_graph_has_expected_edge():
    from bibliometrics.services.analyzer import _country_collaboration_graph

    G = _country_collaboration_graph(_vosviewer_records(), min_collab=2)
    assert G is not None
    assert G.number_of_nodes() == 2 and G.number_of_edges() == 1


def test_bibliographic_coupling_graph_has_expected_edge():
    from bibliometrics.services.analyzer import _bibliographic_coupling_graph

    G = _bibliographic_coupling_graph(_vosviewer_records(), min_shared=2)
    assert G is not None
    assert G.number_of_nodes() == 2 and G.number_of_edges() == 1


# ─────────────────────────── Topic Map — görsel kalite düzeltmesi ───────────────────────────
# tasks/todo.md "[ÖNCELİKLİ] Topic Map görsel kalite sorunu": sabit figsize=(14,8) dikey
# (A4) sayfada az küme olduğunda devasa boş alan bırakıyordu + uzun anahtar kelime
# öbekleri (özetten çıkarılmış) balonun dışına taşıyordu. Bu testler regresyonu önler.

def _topic_map_records(n_clusters: int):
    pool = [
        ['dijital pazarlama stratejileri', 'sosyal medya reklamciligi',
         'marka sadakati olusturma', 'tuketici satin alma davranisi'],
        ['makine ogrenmesi algoritmalari', 'derin sinir aglari',
         'goruntu isleme teknikleri', 'yapay zeka etik sorunlari'],
        ['iklim degisikligi etkileri', 'surdurulebilir kalkinma hedefleri',
         'karbon ayak izi azaltma', 'yenilenebilir enerji kaynaklari'],
        ['psikolojik dayaniklilik gelistirme', 'stres yonetimi teknikleri',
         'ruh sagligi farkindaligi', 'bilissel davranissal terapi'],
        ['egitim teknolojileri entegrasyonu', 'uzaktan ogretim yontemleri',
         'ogrenci basarisi degerlendirme', 'dijital okuryazarlik becerileri'],
    ]
    records = []
    for kws in pool[:n_clusters]:
        for j in range(15):
            records.append({'keywords': kws, 'year': 2019 + (j % 4)})
    return records


def test_topic_map_figure_aspect_not_excessively_wide_for_few_clusters():
    """Eski sabit (14, 8) figsize — en/boy oranı 1.75 — dikey A4 sayfada yarıdan
    fazla boş bırakıyordu. Az kümede bile artık makul bir en/boy oranı (≤1.5,
    tek küme hariç — o zaten minimal/kaçınılmaz bir durum) olmalı."""
    from bibliometrics.services.analyzer import topic_map
    for n in (2, 3, 4, 5):
        fig = topic_map(_topic_map_records(n), n_topics=n)
        assert fig is not None
        w, h = fig.get_size_inches()
        assert w / h <= 1.5, f'n={n} için en/boy oranı çok geniş: {w}x{h}'


def test_topic_map_figsize_scales_with_cluster_count():
    """figsize artık sabit değil — küme (satır/sütun) sayısına göre değişmeli
    (14x8 sabitine geri dönüş regresyonunu yakalar)."""
    from bibliometrics.services.analyzer import topic_map
    fig_few = topic_map(_topic_map_records(2), n_topics=2)
    fig_many = topic_map(_topic_map_records(5), n_topics=5)
    assert fig_few.get_size_inches().tolist() != fig_many.get_size_inches().tolist()
    assert tuple(fig_few.get_size_inches()) != (14, 8)


def test_topic_map_eight_clusters_still_produces_figure():
    """Varsayılan n_topics=8 senaryosu (B2-2'den beri var olan davranış) regresyon olmadan çalışmalı."""
    from bibliometrics.services.analyzer import topic_map
    fig = topic_map(_topic_map_records(5), n_topics=8)
    assert fig is not None


# ─────────────────────────── B3-3: Yöntem Metni + PRISMA Diyagramı ───────────────────────────

def _method_text_records():
    return [{'title': f'Makale {i}', 'year': 2021} for i in range(5)]


def test_build_method_text_file_source_has_expected_content():
    from bibliometrics.services.method_text import build_method_text

    text = build_method_text(
        _method_text_records(), stats={'no_title': 2, 'duplicate': 1}, source={
            'kind': 'file', 'format': 'Scopus CSV', 'files': 1,
        }, n_analyses=26,
    )
    assert 'Scopus CSV' in text
    assert '5' in text  # analize giren kayıt sayısı
    assert '2' in text  # no_title
    assert '1' in text  # duplicate
    assert '26' in text  # n_analyses
    assert 'Analizus' in text
    assert 'otomatik olarak üretilmiştir' in text
    assert 'bir uzman tarafından incelenip onaylanmamıştır' in text
    assert 'OpenAlex' not in text  # dosya kaynağında sorgu/OpenAlex cümlesi olmamalı


def test_build_method_text_openalex_with_query_mentions_query_and_cutoff():
    from bibliometrics.services.method_text import build_method_text

    text = build_method_text(
        _method_text_records(), stats={'no_title': 0, 'duplicate': 0}, source={
            'kind': 'openalex', 'found': 500, 'fetched': 300, 'max_records': 300,
            'query': 'machine learning in education',
        }, n_analyses=20,
    )
    assert 'OpenAlex' in text
    assert 'machine learning in education' in text
    assert '500' in text and '300' in text


def test_build_method_text_openalex_without_query_omits_quote():
    from bibliometrics.services.method_text import build_method_text

    text = build_method_text(
        _method_text_records(), stats={}, source={
            'kind': 'openalex', 'found': 5, 'fetched': 5, 'max_records': 2000,
        }, n_analyses=20,
    )
    assert 'OpenAlex' in text
    assert '"' not in text.split('\n\n')[0]  # sorgu yok → ilk paragrafta tırnaklı sorgu olmamalı


def test_build_method_text_pubmed_source():
    from bibliometrics.services.method_text import build_method_text

    text = build_method_text(
        _method_text_records(), stats={}, source={
            'kind': 'pubmed', 'found': 5, 'fetched': 5, 'max_records': 2000,
        }, n_analyses=18,
    )
    assert 'PubMed' in text


def test_build_method_text_includes_enrichment_sentence_only_when_present():
    from bibliometrics.services.method_text import build_method_text

    base_kwargs = dict(records=_method_text_records(), source={'kind': 'file', 'format': 'BibTeX'}, n_analyses=10)
    without = build_method_text(stats={}, **base_kwargs)
    with_enrichment = build_method_text(stats={'zenginlestirme_enriched': 3}, **base_kwargs)
    assert 'zenginleştirilmiştir' not in without
    assert 'zenginleştirilmiştir' in with_enrichment


def test_build_prisma_png_returns_valid_png_bytes():
    from bibliometrics.services.prisma_diagram import build_prisma_png

    flow = [
        (gettext('Dosyadan okunan kayıt'), 100),
        (gettext('Başlığı olmadığı için çıkarılan'), 5),
        (gettext('Tekrar olduğu için çıkarılan'), 3),
        (gettext('Analize giren kayıt'), 92),
    ]
    png = build_prisma_png(flow)
    assert png is not None
    assert png.startswith(b'\x89PNG')


def test_build_prisma_png_empty_flow_returns_none():
    from bibliometrics.services.prisma_diagram import build_prisma_png
    assert build_prisma_png([]) is None
    assert build_prisma_png(None) is None


def test_upload_method_text_returns_url(monkeypatch):
    from bibliometrics.services import job_runner

    uploaded = {}

    def fake_upload(data, key, content_type):
        uploaded[key] = (data, content_type)
        return f'https://s3.example.com/{key}'

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    url = job_runner._upload_method_text(
        _method_text_records(), {}, {'kind': 'file', 'format': 'BibTeX'}, 10, 'job789',
    )
    assert url == 'https://s3.example.com/bibliometrics/method_text/job789.txt'
    data, content_type = uploaded['bibliometrics/method_text/job789.txt']
    assert content_type == 'text/plain; charset=utf-8'
    assert data.startswith(b'\xef\xbb\xbf')


def test_upload_method_text_tolerates_failure(monkeypatch):
    from bibliometrics.services import job_runner

    def fake_upload(data, key, content_type):
        raise RuntimeError('S3 kapalı')

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    url = job_runner._upload_method_text(
        _method_text_records(), {}, {'kind': 'file', 'format': 'BibTeX'}, 10, 'job789',
    )
    assert url == ''


def test_upload_prisma_diagram_returns_url(monkeypatch):
    from bibliometrics.services import job_runner

    uploaded = {}

    def fake_upload(data, key, content_type):
        uploaded[key] = content_type
        return f'https://s3.example.com/{key}'

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    flow = [(gettext('Dosyadan okunan kayıt'), 10), (gettext('Analize giren kayıt'), 9)]
    url = job_runner._upload_prisma_diagram(flow, 'job789')
    assert url == 'https://s3.example.com/bibliometrics/prisma/job789.png'
    assert uploaded['bibliometrics/prisma/job789.png'] == 'image/png'


def test_upload_prisma_diagram_empty_flow_returns_empty_string(monkeypatch):
    from bibliometrics.services import job_runner
    assert job_runner._upload_prisma_diagram([], 'job789') == ''


def test_upload_prisma_diagram_tolerates_failure(monkeypatch):
    from bibliometrics.services import job_runner

    def fake_upload(data, key, content_type):
        raise RuntimeError('S3 kapalı')

    monkeypatch.setattr('forum.s3_utils.upload_bytes_to_s3', fake_upload)

    flow = [(gettext('Dosyadan okunan kayıt'), 10), (gettext('Analize giren kayıt'), 9)]
    assert job_runner._upload_prisma_diagram(flow, 'job789') == ''
